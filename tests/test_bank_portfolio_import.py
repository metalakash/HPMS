"""The bank workbook importer's parsing, against small workbooks built here in the bank's layout.

No database and no customer data: the borrowers and figures below are invented.
"""

from datetime import date
from decimal import Decimal

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import PatternFill  # noqa: E402

from backend.scripts import import_bank_portfolio as importer  # noqa: E402

GREEN, RED = "FF00B050", "FFFF0000"
QUARTERS = ["Chaitra End", "Ashad End", "Ashoj End", "Poush End"]
# name, MW, CIF, limit, outstanding, remaining, colour, then outstanding / repayment / disbursement per quarter
BORROWERS = [
    ("ALPHA HYDRO PVT LTD", 10, 1001, 500, 400, 100, GREEN, [400, 440, 430, 420], [0, 10, 10, 10], [0, 50, 0, 0]),
    ("BETA POWER LTD", 20, None, 900, None, 900, RED, [0, 0, 300, 600], [0, 0, 0, 0], [0, 0, 300, 300]),
]


def _table(sheet, top, figures_index, title=None):
    if title:
        sheet.cell(top - 1, 4, title)
    for column, label in enumerate(["MW", "S. No.", "CIF", "A/c Name", "Limit Amt",
                                    "Principal Outst as on 31 Chaitra 2081", "Remaining Disbursement"] + QUARTERS, start=1):
        sheet.cell(top, column, label)
    for offset, row in enumerate(BORROWERS, start=1):
        name, mw, cif, limit, outstanding, remaining, colour = row[:7]
        for column, value in enumerate([mw, offset, cif, name, limit, outstanding, remaining] + row[figures_index], start=1):
            sheet.cell(top + offset, column, value)
        sheet.cell(top + offset, 4).fill = PatternFill(fill_type="solid", fgColor=colour)
    total = top + len(BORROWERS) + 1
    sheet.cell(total, 4, "Total of Hydropower Limit Approved")
    for column in range(8, 8 + len(QUARTERS)):
        sheet.cell(total, column, sum(row[figures_index][column - 8] for row in BORROWERS))
    return total


@pytest.fixture
def projection_workbook(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = importer.PORTFOLIO_SHEET
    for row, (label, mw, count, colour) in enumerate(
            [("In Operation", 10, 1, GREEN), ("Yet to Start Drawdown", 20, 1, RED)], start=1):
        sheet.cell(row, 1, mw)
        sheet.cell(row, 2, label)
        sheet.cell(row, 3, count)
        sheet.cell(row, 4).fill = PatternFill(fill_type="solid", fgColor=colour)
    sheet.cell(4, 10, "FY 2082/83")  # only the first full fiscal year is named, as in the bank's sheet
    total = _table(sheet, 5, 7)
    sheet.cell(total, 1, 30)
    sheet.cell(total, 5, 1400)
    sheet.cell(total, 6, 400)
    _table(sheet, 12, 8, title="Repayment")
    _table(sheet, 18, 9, title="Disbursement")
    path = tmp_path / "projection.xlsx"
    book.save(path)
    return path


def test_borrowers_take_their_status_from_the_row_colour(projection_workbook):
    issues = []
    portfolio = importer.read_portfolio(projection_workbook, issues)
    alpha, beta = portfolio["borrowers"]

    assert (alpha["name"], alpha["cif"], alpha["status"]) == ("ALPHA HYDRO PVT LTD", "1001", "in operation")
    assert (beta["cif"], beta["status"], beta["outstanding"]) == (None, "yet to start drawdown", Decimal(0))
    assert portfolio["total"] == {"mw": Decimal(30), "limit": Decimal(1400), "outstanding": Decimal(400)}
    assert portfolio["stated"]["in operation"] == {"mw": Decimal(10), "count": 1}
    assert issues == []


def test_projection_quarters_are_dated_from_the_one_named_fiscal_year(projection_workbook):
    portfolio = importer.read_portfolio(projection_workbook, [])
    quarters = portfolio["quarters"]

    assert [(q["fiscal_year"], q["quarter"], q["period_end_bs"]) for q in quarters] == [
        ("2081/82", 3, "2081-12-31"), ("2081/82", 4, "2082-03-32"),
        ("2082/83", 1, "2082-06-31"), ("2082/83", 2, "2082-09-30"),
    ]
    assert quarters[0]["period_end_ad"] == date(2025, 4, 13)
    assert [q["is_opening"] for q in quarters] == [True, False, False, False]

    alpha = portfolio["borrowers"][0]["projection"]
    assert alpha[1] == {"outstanding": Decimal(440), "repayment": Decimal(10), "disbursement": Decimal(50)}
    # Each quarter is the one before plus disbursement less repayment, and the totals match the sheet's
    for before, after in zip(alpha, alpha[1:]):
        assert before["outstanding"] + after["disbursement"] - after["repayment"] == after["outstanding"]
    assert all(q["outstanding_parsed"] == q["outstanding_stated"] for q in quarters)
    assert quarters[3]["disbursement_parsed"] == Decimal(300)


def test_generation_sheets_read_bs_months_units_and_rate_steps(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    for row in [
        ["Alpha Hydro Pvt. Ltd."], ["Generation Details till date"], ["For Year 2080 BS."],
        ["S. No.", "Month & Year", "PPA Rate (Rs.)", "Contract Energy (CE)", "Contract Revenue (Rs.)",
         "Energy Generated", "Generation Revenue (Rs.)", "% over CE"],
        [1, "Mangshir", 4.8, 1000000, 4800000, 900000, 4320000, 0.9],
        [2, "Poush", 8.4, 500000, 4200000, 400000, 3360000, 0.8],
        [None, "Total", None, 1500000, 9000000, 1300000, 7680000, 0.87],
        ["For Fiscal Year 2081-82 BS."],
        ["S. No.", "Month & Year", "PPA Rate (Rs.)", "Contract Energy (CE)", "Contract Revenue (Rs.)",
         "Energy Generated", "Generation Revenue (Rs.)", "% over CE"],
        [1, "Shrawan -2081", 4.94, 1000000, 4940000, 1000000, 4940000, 1],
        [2, "Baisakh -2082", 4.94, 1000000, 4940000, None, None, None],  # not filled in yet
    ]:
        sheet.append(row)
    path = tmp_path / "generation.xlsx"
    book.save(path)

    issues = []
    plants = importer.read_generation(path, issues)
    months = plants["Alpha Hydro Pvt. Ltd."]
    assert sorted(months) == [(2080, 8), (2080, 9), (2081, 4)]
    assert months[(2080, 9)] == {"contract": Decimal(500000), "generated": Decimal(400000),
                                 "revenue": Decimal(3360000), "rate": Decimal("8.4")}

    borrowers = [{"name": "ALPHA HYDRO PVT LTD"}, {"name": "BETA POWER LTD"}]
    importer.assign_generation(plants, borrowers, issues)
    assert borrowers[0]["plants"] == ["Alpha Hydro Pvt. Ltd."] and borrowers[1]["generation"] == {}
    assert issues == []

    steps = importer.rate_history(borrowers[0]["generation"])
    assert [(s["season"], s["rate"], s["valid_from_bs"], s["valid_to_ad"] is None) for s in steps] == [
        ("wet", Decimal("4.8"), "2080-08-01", False), ("wet", Decimal("4.94"), "2081-04-01", True),
        ("dry", Decimal("8.4"), "2080-09-01", True),
    ]


@pytest.mark.parametrize("label, month", [
    ("Baisakh -2079", 1), ("Jestha-2079", 2), ("Asadh", 3), ("Shrawan ", 4), ("Bhadra", 5), ("Aswin-2080", 6),
    ("Ashwin-2080", 6), ("Kartik", 7), ("Mangsir-2081", 8), ("Mangshir", 8), ("Poush", 9), ("Magh", 10),
    ("Falgun", 11), ("Chaitra", 12), ("Total", None),
])
def test_bs_month_labels(label, month):
    assert importer.bs_month(label) == month
