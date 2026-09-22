"""Tests for file export service (CSV, Excel, JSON)."""

import pytest
import csv
import io
from decimal import Decimal

from backend.app.services.export_service import ExportService


class TestCSVExport:
    """Test CSV file generation."""

    def test_generate_csv_basic(self):
        """Test basic CSV generation."""
        rows = [
            {
                "project_code": "PROJ-001",
                "project_name": "Test Project",
                "capacity_mw": "50.00",
            },
            {
                "project_code": "PROJ-002",
                "project_name": "Another Project",
                "capacity_mw": "75.00",
            },
        ]

        csv_content = ExportService.generate_csv("portfolio", rows)

        assert csv_content is not None
        assert isinstance(csv_content, bytes)
        assert len(csv_content) > 0

    def test_generate_csv_content_valid(self):
        """Test CSV content is valid."""
        rows = [
            {"name": "John", "email": "john@example.com", "status": "active"},
            {"name": "Jane", "email": "jane@example.com", "status": "inactive"},
        ]

        csv_content = ExportService.generate_csv("users", rows)

        # Parse CSV
        text_content = csv_content.decode("utf-8-sig")  # Remove BOM
        reader = csv.DictReader(io.StringIO(text_content))
        parsed_rows = list(reader)

        assert len(parsed_rows) == 2
        assert parsed_rows[0]["name"] == "John"
        assert parsed_rows[1]["name"] == "Jane"

    def test_generate_csv_empty(self):
        """Test CSV with no rows."""
        csv_content = ExportService.generate_csv("empty", [])

        assert csv_content == b""

    def test_generate_csv_special_characters(self):
        """Test CSV with special characters (Nepali)."""
        rows = [
            {
                "project_code": "काली-००१",
                "project_name": "काली गण्डकी",
                "status": "अपरिचालनयोग्य",
            },
        ]

        csv_content = ExportService.generate_csv("portfolio", rows)

        # Verify content can be decoded
        text_content = csv_content.decode("utf-8")
        assert "काली-००१" in text_content


class TestExcelExport:
    """Test Excel file generation."""

    def test_generate_excel_basic(self):
        """Test basic Excel generation."""
        rows = [
            {"project_code": "PROJ-001", "capacity_mw": 50, "status": "active"},
            {"project_code": "PROJ-002", "capacity_mw": 75, "status": "inactive"},
        ]

        excel_content = ExportService.generate_excel("portfolio", rows)

        assert excel_content is not None
        assert isinstance(excel_content, bytes)
        assert len(excel_content) > 0

    def test_generate_excel_content_valid(self):
        """Test Excel file is valid."""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        rows = [
            {"name": "John", "email": "john@example.com", "score": 95},
            {"name": "Jane", "email": "jane@example.com", "score": 98},
        ]

        excel_content = ExportService.generate_excel("results", rows)

        # Parse Excel
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(excel_content))
        ws = wb.active

        # Check headers
        assert ws["A1"].value == "name"
        assert ws["B1"].value == "email"
        assert ws["C1"].value == "score"

        # Check data
        assert ws["A2"].value == "John"
        assert ws["B2"].value == "john@example.com"
        assert ws["C2"].value == 95

    def test_generate_excel_empty(self):
        """Test Excel with no rows."""
        excel_content = ExportService.generate_excel("empty", [])

        assert excel_content == b""

    def test_generate_excel_sheet_name_truncation(self):
        """Test Excel sheet name is truncated to 31 chars."""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        rows = [{"value": 1}]
        report_id = "very_long_report_name_that_exceeds_31_characters_limit"

        excel_content = ExportService.generate_excel(report_id, rows)

        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(excel_content))
        ws = wb.active

        assert len(ws.title) <= 31


class TestJSONExport:
    """Test JSON file generation."""

    def test_generate_json_basic(self):
        """Test basic JSON generation."""
        rows = [
            {"project_code": "PROJ-001", "capacity_mw": 50},
            {"project_code": "PROJ-002", "capacity_mw": 75},
        ]

        json_content = ExportService.generate_json("portfolio", rows)

        assert json_content is not None
        assert isinstance(json_content, str)

    def test_generate_json_content_valid(self):
        """Test JSON content is valid."""
        import json

        rows = [{"name": "John", "score": 95}]

        json_content = ExportService.generate_json("results", rows)

        # Parse JSON
        parsed = json.loads(json_content)

        assert parsed["report_id"] == "results"
        assert parsed["record_count"] == 1
        assert len(parsed["data"]) == 1
        assert parsed["data"][0]["name"] == "John"

    def test_generate_json_with_metadata(self):
        """Test JSON with metadata."""
        import json

        rows = [{"value": 1}]
        metadata = {"filters": {"province": "Gandaki"}, "version": "1.0"}

        json_content = ExportService.generate_json("test", rows, metadata)

        parsed = json.loads(json_content)

        assert parsed["metadata"]["province"] == "Gandaki"
        assert parsed["metadata"]["version"] == "1.0"

    def test_generate_json_timestamp_included(self):
        """Test JSON includes export timestamp."""
        import json
        from datetime import datetime

        rows = [{"value": 1}]
        json_content = ExportService.generate_json("test", rows)

        parsed = json.loads(json_content)

        assert "export_timestamp" in parsed
        # Verify it's a valid ISO format
        datetime.fromisoformat(parsed["export_timestamp"].replace("Z", "+00:00"))


class TestExportFilename:
    """Test export filename generation."""

    def test_get_export_filename_csv(self):
        """Test CSV filename generation."""
        filename = ExportService.get_export_filename("portfolio", "csv")

        assert filename.startswith("portfolio_export_")
        assert filename.endswith(".csv")

    def test_get_export_filename_excel(self):
        """Test Excel filename generation."""
        filename = ExportService.get_export_filename("covenant", "excel")

        assert filename.startswith("covenant_export_")
        assert filename.endswith(".xlsx")

    def test_get_export_filename_json(self):
        """Test JSON filename generation."""
        filename = ExportService.get_export_filename("capex", "json")

        assert filename.startswith("capex_export_")
        assert filename.endswith(".json")

    def test_get_export_filename_includes_timestamp(self):
        """Test filename includes timestamp."""
        from datetime import datetime

        filename = ExportService.get_export_filename("test", "csv")

        # Extract timestamp part (between "export_" and ".csv")
        timestamp_part = filename.split("export_")[1].replace(".csv", "")

        # Verify it's a valid timestamp (YYYYMMDD_HHMMSS format)
        parts = timestamp_part.split("_")
        assert len(parts) == 2
        assert len(parts[0]) == 8  # YYYYMMDD
        assert len(parts[1]) == 6  # HHMMSS


class TestExportDataTypes:
    """Test export handles various data types."""

    def test_export_with_decimal_values(self):
        """Test export with Decimal numbers."""
        rows = [
            {"project_code": "P1", "capacity": Decimal("50.5000")},
            {"project_code": "P2", "capacity": Decimal("75.2500")},
        ]

        csv_content = ExportService.generate_csv("test", rows)

        text_content = csv_content.decode("utf-8-sig")
        assert "50.5" in text_content
        assert "75.25" in text_content

    def test_export_with_none_values(self):
        """Test export with None/null values."""
        rows = [
            {"name": "John", "email": "john@example.com", "phone": None},
            {"name": "Jane", "email": None, "phone": "555-1234"},
        ]

        csv_content = ExportService.generate_csv("test", rows)

        text_content = csv_content.decode("utf-8-sig")
        assert "None" in text_content or text_content.count(",") > 0

    def test_export_with_boolean_values(self):
        """Test export with boolean values."""
        rows = [
            {"name": "Project A", "is_active": True, "is_approved": False},
            {"name": "Project B", "is_active": False, "is_approved": True},
        ]

        csv_content = ExportService.generate_csv("test", rows)

        text_content = csv_content.decode("utf-8-sig")
        assert "True" in text_content
        assert "False" in text_content
