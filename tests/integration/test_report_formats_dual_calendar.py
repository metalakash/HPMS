"""Phase 11.2: Word export, BS+AD dates in every report format, download endpoint."""

import io
import uuid
from datetime import date, datetime
from types import SimpleNamespace

import pytest
from httpx import ASGITransport, AsyncClient

from backend.app.database import get_db
from backend.app.main import app
from backend.app.security.auth_middleware import CurrentUser, get_current_user
from backend.app.services.export_service import ExportService
from backend.app.services.report_dates import add_bs_columns, bs_string, dual_stamp

ROWS = [
    {"project_code": "P1", "capacity_mw": 42.5, "created_date_ad": "2026-04-14", "created_date_bs": None},
    {"project_code": "P2", "capacity_mw": 10.0, "created_date_ad": None, "created_date_bs": None},
]


def test_add_bs_columns_fills_from_ad_and_keeps_existing():
    rows = add_bs_columns([dict(r) for r in ROWS])
    assert rows[0]["created_date_bs"] == "2083-01-01"
    assert rows[1]["created_date_bs"] is None  # no AD date, nothing to convert
    kept = add_bs_columns([{"x_ad": "2026-04-14", "x_bs": "custom"}])
    assert kept[0]["x_bs"] == "custom"


def test_add_bs_columns_handles_date_objects_and_out_of_range():
    rows = add_bs_columns([{"d_ad": date(2026, 4, 14), "d_bs": None}, {"d_ad": "1800-01-01", "d_bs": None}])
    assert rows[0]["d_bs"] == "2083-01-01"
    assert rows[1]["d_bs"] is None


def test_dual_stamp_has_both_calendars():
    assert dual_stamp(datetime(2026, 4, 14, 3, 0)) == "2026-04-14 AD / 2083-01-01 BS"


def test_word_export_contains_table_stamp_and_filters():
    from docx import Document
    data = ExportService.generate_word("portfolio", add_bs_columns([dict(r) for r in ROWS]),
                                       filters={"province": "Gandaki"})
    assert data[:2] == b"PK"
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "AD /" in text and "BS" in text
    assert "province=Gandaki" in text
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells][:2] == ["project_code", "capacity_mw"]
    assert table.rows[1].cells[0].text == "P1"
    assert "2083-01-01" in [c.text for c in table.rows[1].cells]
    assert table.rows[2].cells[-1].text == ""  # None renders blank, not "None"


def test_word_export_empty_returns_empty_bytes():
    assert ExportService.generate_word("portfolio", []) == b""


def test_word_filename_extension():
    assert ExportService.get_export_filename("portfolio", "word").endswith(".docx")


def test_excel_has_report_info_sheet_with_dual_stamp():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(ExportService.generate_excel("portfolio", ROWS)))
    assert wb.sheetnames[0] == "portfolio" and "Report info" in wb.sheetnames
    assert "BS" in wb["Report info"]["B2"].value


def test_pdf_header_uses_dual_stamp(monkeypatch):
    import reportlab.platypus as platypus
    from backend.app.services import pdf_service
    seen = []
    real = platypus.Paragraph
    monkeypatch.setattr(platypus, "Paragraph", lambda text, *a, **k: (seen.append(text), real(text, *a, **k))[1])
    pdf = pdf_service.PDFService.generate_portfolio_report(
        {"total_projects": 0, "total_capacity_mw": 0, "active_projects": 0, "average_rate": None, "projects": []})
    assert pdf[:4] == b"%PDF"
    assert any("Generated:" in t and " BS" in t for t in seen)


def test_routes_reports_imports_rls_service():
    import backend.app.api.routes_reports as m
    assert m.RLSService is not None  # covenant PDF referenced it without importing (NameError -> 500)


@pytest.fixture
async def api(monkeypatch):
    async def fake_db():
        yield SimpleNamespace()

    async def fake_loader(db, request, user_id):
        return [dict(r) for r in ROWS], len(ROWS)

    monkeypatch.setattr("backend.app.api.routes_reports._load_rows", fake_loader)
    app.dependency_overrides[get_db] = fake_db
    app.dependency_overrides[get_current_user] = lambda: CurrentUser({
        "sub": str(uuid.uuid4()), "username": "tester", "roles": ["admin"], "is_authenticated": True})
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://localhost") as client:
        yield client
    app.dependency_overrides.clear()


@pytest.mark.parametrize("fmt,ctype,magic", [
    ("word", "wordprocessingml.document", b"PK"),
    ("excel", "spreadsheetml.sheet", b"PK"),
    ("csv", "text/csv", b"\xef\xbb\xbf"),
])
async def test_download_endpoint_returns_file(api, fmt, ctype, magic):
    r = await api.post("/api/v1/reports/export/download", json={"report_id": "portfolio", "format": fmt})
    assert r.status_code == 200, r.text
    assert ctype in r.headers["content-type"]
    assert r.content.startswith(magic)
    assert "attachment" in r.headers["content-disposition"]


async def test_download_rejects_json_format(api):
    r = await api.post("/api/v1/reports/export/download", json={"report_id": "portfolio", "format": "json"})
    assert r.status_code == 400


async def test_download_requires_portfolio_role(api):
    app.dependency_overrides[get_current_user] = lambda: CurrentUser({
        "sub": str(uuid.uuid4()), "username": "g", "roles": ["guest"], "is_authenticated": True})
    r = await api.post("/api/v1/reports/export/download", json={"report_id": "portfolio", "format": "word"})
    assert r.status_code == 403
