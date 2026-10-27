"""Small workbooks in the layout of the bank's two portfolio workbooks, with invented borrowers and figures.

Used by the importer's parsing tests and by the in-app import tests. No customer data.
"""

import io

import openpyxl
from openpyxl.styles import PatternFill

from backend.scripts import import_bank_portfolio as importer

GREEN, RED = "FF00B050", "FFFF0000"
QUARTERS = ["Chaitra End", "Ashad End", "Ashoj End", "Poush End"]
# name, MW, CIF, limit, outstanding, remaining, colour, then outstanding / repayment / disbursement per quarter
BORROWERS = [
    ("ALPHA HYDRO PVT LTD", 10, 1001, 500, 400, 100, GREEN, [400, 440, 430, 420], [0, 10, 10, 10], [0, 50, 0, 0]),
    ("BETA POWER LTD", 20, None, 900, None, 900, RED, [0, 0, 300, 600], [0, 0, 0, 0], [0, 0, 300, 300]),
]


def _table(sheet, top, figures_index, title=None):
    if title:
        sheet.cell(top - 1, 4, title)
    for column, label in enumerate(["MW", "S. No.", "CIF", "A/c Name", "Limit Amt",
                                    "Principal Outst as on 31 Chaitra 2081", "Remaining Disbursement"] + QUARTERS, start=1):
        sheet.cell(top, column, label)
    for offset, row in enumerate(BORROWERS, start=1):
        name, mw, cif, limit, outstanding, remaining, colour = row[:7]
        for column, value in enumerate([mw, offset, cif, name, limit, outstanding, remaining] + row[figures_index], start=1):
            sheet.cell(top + offset, column, value)
        sheet.cell(top + offset, 4).fill = PatternFill(fill_type="solid", fgColor=colour)
    total = top + len(BORROWERS) + 1
    sheet.cell(total, 4, "Total of Hydropower Limit Approved")
    for column in range(8, 8 + len(QUARTERS)):
        sheet.cell(total, column, sum(row[figures_index][column - 8] for row in BORROWERS))
    return total


def projection_book(total_limit: int = 1400) -> openpyxl.Workbook:
    """The loan projection workbook: legend, borrower table and its three projection tables.
    A ``total_limit`` other than 1400 makes the stated total disagree with the rows."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = importer.PORTFOLIO_SHEET
    for row, (label, mw, count, colour) in enumerate(
            [("In Operation", 10, 1, GREEN), ("Yet to Start Drawdown", 20, 1, RED)], start=1):
        sheet.cell(row, 1, mw)
        sheet.cell(row, 2, label)
        sheet.cell(row, 3, count)
        sheet.cell(row, 4).fill = PatternFill(fill_type="solid", fgColor=colour)
    sheet.cell(4, 10, "FY 2082/83")  # only the first full fiscal year is named, as in the bank's sheet
    total = _table(sheet, 5, 7)
    sheet.cell(total, 1, 30)
    sheet.cell(total, 5, total_limit)
    sheet.cell(total, 6, 400)
    _table(sheet, 12, 8, title="Repayment")
    _table(sheet, 18, 9, title="Disbursement")
    return book


GENERATION_ROWS = [
    ["Alpha Hydro Pvt. Ltd."], ["Generation Details till date"], ["For Year 2080 BS."],
    ["S. No.", "Month & Year", "PPA Rate (Rs.)", "Contract Energy (CE)", "Contract Revenue (Rs.)",
     "Energy Generated", "Generation Revenue (Rs.)", "% over CE"],
    [1, "Mangshir", 4.8, 1000000, 4800000, 900000, 4320000, 0.9],
    [2, "Poush", 8.4, 500000, 4200000, 400000, 3360000, 0.8],
    [None, "Total", None, 1500000, 9000000, 1300000, 7680000, 0.87],
    ["For Fiscal Year 2081-82 BS."],
    ["S. No.", "Month & Year", "PPA Rate (Rs.)", "Contract Energy (CE)", "Contract Revenue (Rs.)",
     "Energy Generated", "Generation Revenue (Rs.)", "% over CE"],
    [1, "Shrawan -2081", 4.94, 1000000, 4940000, 1000000, 4940000, 1],
    [2, "Baisakh -2082", 4.94, 1000000, 4940000, None, None, None],  # not filled in yet
]


def generation_book() -> openpyxl.Workbook:
    book = openpyxl.Workbook()
    sheet = book.active
    for row in GENERATION_ROWS:
        sheet.append(row)
    return book


def as_bytes(book: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
