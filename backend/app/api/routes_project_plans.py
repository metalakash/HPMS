"""A project's budget against actual cost (RFP FUNC B.2, C.4, F.9) and its quarterly loan projection (B.6).

Under /api/v1/projects/{project_id}/. Read-only, plus a preview of a budget upload: both change through
change requests (BUDGET_LINE, LOAN_PROJECTION; see services/budget_changes.py and projection_changes.py).
"""
import csv
import io
import uuid
from decimal import Decimal
from typing import Any, Dict, List

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.api.routes_projects import _get_visible_project
from backend.app.database import get_db
from backend.app.models.financial import BudgetLine, LoanProjectionQuarter
from backend.app.security.auth_middleware import CurrentUser, get_current_user
from backend.app.services.budget_changes import MAX_LINES, line_view, variance
from backend.app.services.change_kinds import amount

router = APIRouter(prefix="/api/v1/projects/{project_id}", tags=["project-plans"])


@router.get("/budget")
async def project_budget(project_id: str, db: AsyncSession = Depends(get_db),
                         user: CurrentUser = Depends(get_current_user)):
    project = await _get_visible_project(db, user, project_id)
    lines = (await db.execute(select(BudgetLine).where(BudgetLine.project_id == project.id)
                              .order_by(BudgetLine.category))).scalars().all()
    budgeted = sum((line.budgeted_amount or Decimal("0") for line in lines), Decimal("0"))
    actual = sum((line.actual_amount or Decimal("0") for line in lines), Decimal("0"))
    total_variance, total_pct = variance(budgeted, actual) if lines else (None, None)
    return {"data": {
        "lines": [line_view(line) for line in lines],
        "totals": {"budgeted_amount": str(budgeted), "actual_amount": str(actual),
                   "variance_amount": None if total_variance is None else str(total_variance),
                   "variance_pct": None if total_pct is None else str(total_pct),
                   "overspent_lines": sum(1 for line in lines if (line.variance_amount or 0) > 0)},
    }}


# Column headings accepted for each field, compared without case, spaces or underscores
HEADINGS = {
    "category": {"category", "costcategory", "head", "budgethead", "item"},
    "budgeted_amount": {"budgetedamount", "budget", "budgeted", "budgetnpr", "budgetedamountnpr"},
    "actual_amount": {"actualamount", "actual", "actualcost", "spent", "actualnpr", "actualamountnpr"},
}
MAX_UPLOAD_BYTES = 2 * 1024 * 1024


def _key(heading: Any) -> str:
    return "".join(ch for ch in str(heading or "").lower() if ch.isalnum())


def _rows(name: str, content: bytes) -> List[List[Any]]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    if name.lower().endswith(".csv"):
        return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    raise HTTPException(status_code=422, detail="Upload an Excel (.xlsx) or CSV file")


@router.post("/budget/preview")
async def preview_budget_upload(project_id: str, file: UploadFile = File(...), db: AsyncSession = Depends(get_db),
                                user: CurrentUser = Depends(get_current_user)):
    """Read a budget sheet and say what it would change. Writes nothing: the lines are then proposed as one
    change request. A category already on the budget updates that line; a new one adds a line."""
    project = await _get_visible_project(db, user, project_id)
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=422, detail="The file is larger than 2 MB")
    try:
        rows = [row for row in _rows(file.filename or "", content) if any(str(c or "").strip() for c in row)]
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="The file could not be read as a spreadsheet")
    if not rows:
        raise HTTPException(status_code=422, detail="The file is empty")

    columns: Dict[str, int] = {}
    for index, heading in enumerate(rows[0]):
        for field, names in HEADINGS.items():
            if _key(heading) in names and field not in columns:
                columns[field] = index
    missing = [field for field in ("category", "budgeted_amount") if field not in columns]
    if missing:
        raise HTTPException(status_code=422, detail=(
            f"No {' or '.join(m.replace('_', ' ') for m in missing)} column. The first row must name the columns: "
            "Category, Budgeted amount, Actual amount"))
    if len(rows) - 1 > MAX_LINES:
        raise HTTPException(status_code=422, detail=f"At most {MAX_LINES} lines in one upload")

    existing = {(line.category or "").strip().lower(): line for line in (await db.execute(
        select(BudgetLine).where(BudgetLine.project_id == project.id))).scalars()}
    lines, errors, seen = [], [], set()
    for number, row in enumerate(rows[1:], start=2):
        def cell(field):
            index = columns.get(field)
            return row[index] if index is not None and index < len(row) else None

        category = str(cell("category") or "").strip()
        try:
            if not category:
                raise ValueError("no category")
            if len(category) > 100:
                raise ValueError("category is longer than 100 characters")
            if category.lower() in seen:
                raise ValueError(f"{category} is listed twice")
            budgeted = amount(cell("budgeted_amount"), "budgeted amount") if cell("budgeted_amount") not in (None, "") else None
            if budgeted is None:
                raise ValueError("no budgeted amount")
            actual = amount(cell("actual_amount") or 0, "actual amount")
        except ValueError as problem:
            errors.append({"row": number, "message": str(problem)})
            continue
        seen.add(category.lower())
        match = existing.get(category.lower())
        difference, pct = variance(budgeted, actual)
        lines.append({"id": str(match.id) if match else None, "category": category, "budgeted_amount": str(budgeted),
                      "actual_amount": str(actual), "variance_amount": str(difference),
                      "variance_pct": None if pct is None else str(pct), "change": "update" if match else "add"})
    return {"data": {"lines": lines, "errors": errors, "upload_batch_id": f"UPLOAD-{uuid.uuid4().hex[:12]}",
                     "file_name": file.filename}}


@router.get("/projection")
async def project_projection(project_id: str, db: AsyncSession = Depends(get_db),
                             user: CurrentUser = Depends(get_current_user)):
    """The project's loan projection by fiscal quarter, opening position first."""
    project = await _get_visible_project(db, user, project_id)
    rows = (await db.execute(select(LoanProjectionQuarter).where(LoanProjectionQuarter.project_id == project.id)
                             .order_by(LoanProjectionQuarter.period_end_ad))).scalars().all()
    return {"data": [{
        "fiscal_year": row.fiscal_year, "quarter": row.quarter, "period_end_ad": row.period_end_ad.isoformat(),
        "period_end_bs": row.period_end_bs, "is_opening": bool(row.is_opening),
        "projected_disbursement": str(row.projected_disbursement), "projected_repayment": str(row.projected_repayment),
        "projected_outstanding": str(row.projected_outstanding), "source_reference": row.source_reference,
    } for row in rows]}
