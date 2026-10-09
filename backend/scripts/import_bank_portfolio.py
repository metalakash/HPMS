"""Load the bank's own hydropower book from its two working spreadsheets into a LOCAL database.

Reads
  * the projection workbook (sheet ``Disbursement_Ashad_2082``): one row per borrower with MW, CIF,
    limit, principal outstanding and remaining disbursement; the status is the row's fill colour.
    The same sheet carries the quarterly projection in three tables: outstanding, repayment, disbursement.
    Sheet ``Energy Bond`` lists the bonds held, and ``Net Loan with existing limit`` the bank's total
    loans and the regulator's minimum share for energy, quarter by quarter. ``New Loan Projection``
    and ``Net Loan after New Loan`` hold the limits the bank plans to approve and their drawdown.
  * the generation workbook: one sheet per plant, monthly contract energy, generation and PPA rate.

and writes projects, loan accounts, the quarterly projection, the bond register, the inputs to the
energy-financing ratio, monthly generation and the PPA rate history. The ratio itself is not copied:
the application works it out, and the report shows where that differs from the workbook's own figure. Nothing is invented:
a field the workbooks do not carry (interest rate, maturity, COD, covenants) is left empty. Province
and district are looked up in the DoED licence list by promoter name where that matches.

The workbooks hold customer data. Keep them outside the repository; this script refuses to write to
anything but a database on this machine. Without ``--apply`` it only prints the reconciliation.

WARNING: ``--apply`` deletes every project and everything that hangs off one, like
seed_realistic_data does. Run that script again to go back to the synthetic portfolio.

Run from the repository root:
    python -m backend.scripts.import_bank_portfolio --projection <xlsx> --generation <xlsx> [--apply]
"""

import argparse
import asyncio
import csv
import re
import sys
import uuid
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import nepali_datetime

MONEY = Decimal("0.01")
PORTFOLIO_SHEET = "Disbursement_Ashad_2082"
BOND_SHEET = "Energy Bond"
RATIO_SHEET = "Net Loan with existing limit"
PIPELINE_SHEET = "New Loan Projection"
PIPELINE_RATIO_SHEET = "Net Loan after New Loan"
DOED_SOURCE = Path(__file__).resolve().parent.parent / "data" / "merged_hydropower_master.csv"
PROVENANCE = "BANK_WORKBOOK"

# Workbook status label -> (pipeline_status, project_stage)
STATUSES = {
    "in operation": ("under_operation", "operation"),
    "under construction": ("under_construction", "construction"),
    "yet to start drawdown": ("yet_to_start_drawdown", "construction"),
    "settled": ("settled", "operation"),
    "proposal under pipeline": ("proposal_under_pipeline", "feasibility"),
}
# First letters of a BS month label -> month number. Ashwin is tested before Ashadh.
MONTH_PREFIXES = [
    (("asw", "ashw", "aso", "asho"), 6), (("bai",), 1), (("jes", "jet"), 2), (("asa", "asha"), 3),
    (("shr", "sra", "sau"), 4), (("bha",), 5), (("kar",), 7), (("man",), 8), (("pou", "pau", "pus"), 9),
    (("mag",), 10), (("fal", "pha", "fag"), 11), (("cha",), 12),
]
# Quarter-end label -> (quarter of the fiscal year, BS month it ends with). Ashad falls in the fiscal year's second year.
QUARTER_ENDS = {"ashoj": (1, 6), "poush": (2, 9), "chaitra": (3, 12), "ashad": (4, 3)}
QUARTER_END_MONTHS = dict(QUARTER_ENDS.values())
DRY_MONTHS = {9, 10, 11, 12}  # Poush to Chaitra: the months the workbooks price at the dry-season rate
NAME_NOISE = {"pvt", "private", "ltd", "limited", "p", "l", "li", "co", "company", "c", "dev", "development"}


def dec(value: Any) -> Optional[Decimal]:
    if value is None or isinstance(value, str):
        return None
    return Decimal(str(value))


def money(value: Optional[Decimal]) -> Optional[Decimal]:
    return None if value is None else value.quantize(MONEY)


def name_key(name: str) -> str:
    """A company name reduced to the words that identify it, so spelling variants compare equal."""
    text = name.lower().replace("hidro", "hydro").replace("hydro power", "hydropower").replace("&", " ")
    text = re.sub(r"\(.*?\)", " ", text)
    return " ".join(w for w in re.sub(r"[^a-z0-9 ]", " ", text).split() if w not in NAME_NOISE)


def bs_month(label: Any) -> Optional[int]:
    letters = re.sub(r"[^a-z]", "", str(label).lower())
    for prefixes, month in MONTH_PREFIXES:
        if letters.startswith(prefixes):
            return month
    return None


def bs_month_start(year: int, month: int) -> date:
    return nepali_datetime.date(year, month, 1).to_datetime_date()


def bs_month_end(year: int, month: int) -> Tuple[date, str]:
    """Last day of a BS month, as an AD date and a BS string."""
    last_ad = bs_month_start(year + month // 12, month % 12 + 1) - timedelta(days=1)
    last = nepali_datetime.date.from_datetime_date(last_ad)
    return last_ad, f"{last.year:04d}-{last.month:02d}-{last.day:02d}"


# ---------------------------------------------------------------------------------------------
# Projection workbook
# ---------------------------------------------------------------------------------------------

def _fill(cell) -> Optional[str]:
    fill = cell.fill
    if fill is None or fill.fill_type is None:
        return None
    colour = fill.fgColor
    if colour.type == "theme":
        return f"theme{colour.theme}/{colour.tint:.2f}"
    return None if colour.rgb in ("FFFFFFFF", "00000000") else str(colour.rgb)


def read_portfolio(path: Path, issues: List[str]) -> Dict[str, Any]:
    """Borrower rows of the projection workbook, plus the totals the workbook states for itself."""
    import openpyxl

    sheet = openpyxl.load_workbook(path, data_only=True)[PORTFOLIO_SHEET]
    header_row = next(r for r in range(1, sheet.max_row + 1)
                      if str(sheet.cell(r, 4).value).strip() == "A/c Name" and sheet.cell(r, 1).value == "MW")

    # The block above the table is the legend: status label, its colour, and the bank's own subtotals
    legend: Dict[str, str] = {}
    stated: Dict[str, Dict[str, Any]] = {}
    for r in range(1, header_row):
        label = str(sheet.cell(r, 2).value or "").strip().lower()
        if label in STATUSES:
            stated[label] = {"mw": dec(sheet.cell(r, 1).value) or Decimal(0), "count": sheet.cell(r, 3).value or 0}
            for c in range(1, 5):
                if _fill(sheet.cell(r, c)):
                    legend[_fill(sheet.cell(r, c))] = label

    borrowers: List[Dict[str, Any]] = []
    total: Dict[str, Any] = {}
    bond_holding: Optional[Decimal] = None
    for r in range(header_row + 1, sheet.max_row + 1):
        name = sheet.cell(r, 4).value
        if not name:
            continue
        name = " ".join(str(name).split())
        limit, outstanding = dec(sheet.cell(r, 5).value), dec(sheet.cell(r, 6).value)
        if name.lower().startswith("total of hydropower"):
            total = {"mw": dec(sheet.cell(r, 1).value), "limit": limit, "outstanding": outstanding}
            bond_row = next((x for x in range(r + 1, r + 4)
                             if str(sheet.cell(x, 4).value or "").lower().startswith("investment in energy bond")), None)
            bond_holding = dec(sheet.cell(bond_row, 5).value) if bond_row else None
            break
        fills = {_fill(sheet.cell(r, c)) for c in range(1, 8)} - {None}
        status = next((legend[f] for f in fills if f in legend), None)
        cif = sheet.cell(r, 3).value
        borrowers.append({
            "offset": r - header_row, "row": r, "name": name, "mw": dec(sheet.cell(r, 1).value), "serial": sheet.cell(r, 2).value,
            "cif": str(cif).strip() if cif else None, "limit": limit or Decimal(0),
            "outstanding": outstanding or Decimal(0), "remaining": dec(sheet.cell(r, 7).value) or Decimal(0),
            "status": status,
        })
    quarters = read_projection(sheet, header_row, borrowers, issues)
    return {"borrowers": borrowers, "total": total, "stated": stated, "quarters": quarters,
            "bond_holding": bond_holding,
            "as_of_label": str(sheet.cell(header_row, 6).value).strip()}


def quarter_columns(sheet, label_row: int, first_column: int) -> List[Tuple[int, int]]:
    """(column, quarter of the fiscal year) for the run of quarter-end headings starting at a column."""
    labelled: List[Tuple[int, int]] = []
    column = first_column
    while True:
        label = str(sheet.cell(label_row, column).value or "").split()
        if not label or label[0].lower() not in QUARTER_ENDS:
            return labelled
        labelled.append((column, QUARTER_ENDS[label[0].lower()][0]))
        column += 1


def date_quarters(sheet, year_row: int, labelled: List[Tuple[int, int]], what: str, issues: List[str]) -> List[Dict[str, Any]]:
    """Fiscal year and end date for each quarter column.

    Only some columns name their fiscal year, so quarters are counted outwards from the first that does.
    """
    anchor = None
    for index, (c, quarter) in enumerate(labelled):
        named = re.search(r"(\d{4})/", str(sheet.cell(year_row, c).value or ""))
        if named:
            anchor = int(named.group(1)) * 4 + quarter - 1 - index
            break
    if anchor is None:
        issues.append(f"{what}: no fiscal year heading found; not loaded")
        return []
    quarters: List[Dict[str, Any]] = []
    for index, (c, quarter) in enumerate(labelled):
        start_year, expected = divmod(anchor + index, 4)
        if expected + 1 != quarter:
            issues.append(f"{what}: column {c} is out of quarter sequence; it stops before that column")
            break
        end_ad, end_bs = bs_month_end(start_year + (1 if quarter == 4 else 0), QUARTER_END_MONTHS[quarter])
        quarters.append({"column": c, "fiscal_year": f"{start_year}/{(start_year + 1) % 100:02d}", "quarter": quarter,
                         "period_end_ad": end_ad, "period_end_bs": end_bs})
    return quarters


def read_projection(sheet, header_row: int, borrowers: List[Dict[str, Any]], issues: List[str]) -> List[Dict[str, Any]]:
    """The quarter columns of the sheet and, on each borrower, its figures for them.

    The sheet repeats the borrower table three times: projected outstanding (the table the
    borrowers were read from), then repayment, then disbursement, with the same rows and columns.
    """
    for b in borrowers:
        b["projection"] = []
    blocks = [r for r in range(header_row, sheet.max_row + 1) if str(sheet.cell(r, 4).value).strip() == "A/c Name"]
    titles = {str(sheet.cell(r - above, 4).value or "").strip().lower(): r for r in blocks for above in (1, 2)}
    if len(blocks) != 3 or "repayment" not in titles or "disbursement" not in titles:
        issues.append("Projection: the repayment and disbursement tables were not found; no projection loaded")
        return []
    tables = {"outstanding": header_row, "repayment": titles["repayment"], "disbursement": titles["disbursement"]}

    # Quarter columns start after 'Remaining Disbursement' and run for as long as they are labelled
    labelled = quarter_columns(sheet, header_row, 8)
    column = labelled[-1][0] + 1 if labelled else 8
    unlabelled = sum(1 for c in range(column, sheet.max_column + 1)
                     if any(sheet.cell(header_row + b["offset"], c).value for b in borrowers))
    if unlabelled:
        issues.append(f"Projection: {unlabelled} further columns hold figures but have no quarter heading; not loaded")

    quarters = date_quarters(sheet, header_row - 1, labelled, "Projection", issues)
    for index, quarter in enumerate(quarters):
        quarter["is_opening"] = index == 0

    def figure(row: int, col: int) -> Decimal:
        return dec(sheet.cell(row, col).value) or Decimal(0)

    for b in borrowers:
        if any(" ".join(str(sheet.cell(row + b["offset"], 4).value or "").split()) != b["name"] for row in tables.values()):
            issues.append(f"{b['name']}: the three projection tables do not line up on its row; no projection")
            continue
        rows = [{name: figure(row + b["offset"], q["column"]) for name, row in tables.items()} for q in quarters]
        if any(any(r.values()) for r in rows):
            b["projection"] = rows

    for name, row in tables.items():
        total_row = next((r for r in range(row + 1, sheet.max_row + 1)
                          if str(sheet.cell(r, 4).value or "").lower().startswith("total of hydropower")), None)
        for index, q in enumerate(quarters):
            q[f"{name}_parsed"] = sum((b["projection"][index][name] for b in borrowers if b["projection"]), Decimal(0))
            q[f"{name}_stated"] = figure(total_row, q["column"]) if total_row else None
    return quarters


def _amount(value: Any) -> Optional[Decimal]:
    """A figure that may have been typed as text, e.g. '  500,000,000.00  '."""
    if isinstance(value, str):
        text = value.replace(",", "").strip()
        return Decimal(text) if re.fullmatch(r"-?\d+(\.\d+)?", text) else None
    return dec(value)


def read_bonds(book, held_total: Optional[Decimal], issues: List[str]) -> List[Dict[str, Any]]:
    """The bond register. The sheet keeps more than one list; the one that adds up to the holding
    the projection sheet states is the current one."""
    if BOND_SHEET not in book.sheetnames:
        issues.append(f"Bonds: no '{BOND_SHEET}' sheet; no bonds loaded")
        return []
    rows = list(book[BOND_SHEET].iter_rows(values_only=True))
    tables: List[Tuple[List[Dict[str, Any]], List[str]]] = []  # each list with the notes raised reading it
    for number, row in enumerate(rows):
        headers = [str(c).strip().lower() if c else "" for c in row]
        if not any("bond" in h for h in headers) or not any("allotted" in h for h in headers):
            continue
        at = {key: next((i for i, h in enumerate(headers) if h.startswith(start)), None)
              for key, start in (("name", "government bond"), ("invested", "inv"), ("amount", "allotted"),
                                 ("yield", "yield"), ("maturity", "maturity"))}
        table, notes = [], []
        for entry in rows[number + 1:]:
            name = entry[at["name"]] if at["name"] is not None else None
            amount = _amount(entry[at["amount"]])
            if not name or amount is None:
                break
            invested = entry[at["invested"]] if at["invested"] is not None else None
            matures = entry[at["maturity"]] if at["maturity"] is not None else None
            invested = invested.date() if hasattr(invested, "date") else None
            matures = matures.date() if hasattr(matures, "date") else None
            if invested and matures and matures < invested:
                # A two-digit year read into the wrong century, e.g. 1933 for 2033
                matures = matures.replace(year=matures.year + 100)
                notes.append(f"Bond '{' '.join(str(name).split())}': maturity is before the investment date in the "
                             f"sheet; read as {matures.isoformat()}")
            rate = _amount(entry[at["yield"]]) if at["yield"] is not None else None
            table.append({"name": " ".join(str(name).split()), "amount": amount, "investment_date": invested,
                          "maturity_date": matures, "yield_pct": rate * 100 if rate is not None else None})
        if table:
            tables.append((table, notes))
    if not tables:
        issues.append("Bonds: no bond table found; no bonds loaded")
        return []
    matching = [t for t in tables if held_total is not None and abs(sum(b["amount"] for b in t[0]) - held_total) < 1]
    if not matching:
        issues.append("Bonds: no list on the bond sheet adds up to the holding the projection sheet states; the last list is used")
    table, notes = (matching or tables)[-1]
    issues.extend(notes)
    return table


def read_ratio_inputs(book, opening: Optional[date], issues: List[str]) -> List[Dict[str, Any]]:
    """Per quarter: the bank's total loans, the regulator's minimum share, and for quarters up to the
    projection's opening date the hydropower outstanding and bond holding the bank recorded."""
    if RATIO_SHEET not in book.sheetnames:
        issues.append(f"Energy financing: no '{RATIO_SHEET}' sheet; ratio inputs not loaded")
        return []
    sheet = book[RATIO_SHEET]
    rows: Dict[str, int] = {}
    for r in range(1, sheet.max_row + 1):
        label = str(sheet.cell(r, 1).value or "").strip().lower()
        for key, start in (("required", "net working"), ("years", "details"), ("hydro", "hydropower loan o/s"),
                           ("bonds", "add: investement in energy bond"), ("loans", "estimated bank total loan"),
                           ("ratio", "% on energy financing on existing")):
            if label.startswith(start):
                rows.setdefault(key, r)
    if {"required", "years", "hydro", "bonds", "loans"} - rows.keys():
        issues.append("Energy financing: the ratio sheet's rows were not recognised; ratio inputs not loaded")
        return []
    quarters = date_quarters(sheet, rows["years"], quarter_columns(sheet, rows["years"] + 1, 2), "Energy financing", issues)
    for q in quarters:
        c = q.pop("column")
        recorded = opening is not None and q["period_end_ad"] <= opening
        share = dec(sheet.cell(rows["required"], c).value)
        stated = dec(sheet.cell(rows["ratio"], c).value) if "ratio" in rows else None
        q.update({
            "bank_total_loans": dec(sheet.cell(rows["loans"], c).value) or None,
            "required_pct": share * 100 if share else None,
            "hydro_outstanding_actual": dec(sheet.cell(rows["hydro"], c).value) if recorded else None,
            "energy_bond_actual": dec(sheet.cell(rows["bonds"], c).value) if recorded else None,
            # past the workbook's last bank-loan estimate its ratio cells hold other formulas, not a share
            "stated_ratio_pct": stated * 100 if stated and stated < 1 else None,
        })
    # Nothing can be tested more than two quarters after the last bank-loan figure
    last = max((i for i, q in enumerate(quarters) if q["bank_total_loans"]), default=-1)
    return quarters[:last + 3]


def read_pipeline(book, issues: List[str]) -> Dict[str, Any]:
    """New limits the bank plans to approve, and the disbursement the workbook schedules from them."""
    limits: List[Dict[str, Any]] = []
    if PIPELINE_SHEET in book.sheetnames:
        rows = list(book[PIPELINE_SHEET].iter_rows(values_only=True))
        profile: List[Decimal] = []
        for number, row in enumerate(rows[:-1]):
            if any(str(c).strip().lower() == "1st year" for c in row if c):
                shares = [dec(c) for c in rows[number + 1][1:]]
                profile = [share * 100 for share in shares if share is not None and 0 < share < 1]
                break
        header = next((i for i, row in enumerate(rows) if str(row[0]).strip().lower() == "year"
                       and str(row[1]).strip().lower() == "new limit"), None)
        for row in rows[header + 1:] if header is not None else []:
            if not re.fullmatch(r"\d{4}/\d{2}", str(row[0]).strip()):
                break
            if dec(row[1]):
                limits.append({"fiscal_year": str(row[0]).strip(), "new_limit": dec(row[1]),
                               "drawdown_pct": [float(share) for share in profile] or None})
    else:
        issues.append(f"New loans: no '{PIPELINE_SHEET}' sheet; planned limits not loaded")

    quarters: List[Dict[str, Any]] = []
    stated: Dict[date, Decimal] = {}
    if PIPELINE_RATIO_SHEET in book.sheetnames:
        sheet = book[PIPELINE_RATIO_SHEET]
        rows_at: Dict[str, int] = {}
        for r in range(1, sheet.max_row + 1):
            label = str(sheet.cell(r, 1).value or "").strip().lower()
            for key, start in (("years", "details"), ("ratio", "% on energy financing with"),
                               ("new", "add: estimated disbursement from required new loan")):
                if label.startswith(start):
                    rows_at.setdefault(key, r)
        if {"years", "new"} <= rows_at.keys():
            for q in date_quarters(sheet, rows_at["years"], quarter_columns(sheet, rows_at["years"] + 1, 2), "New loans", issues):
                c = q.pop("column")
                ratio = dec(sheet.cell(rows_at["ratio"], c).value) if "ratio" in rows_at else None
                if ratio and ratio < 1:
                    stated[q["period_end_ad"]] = ratio * 100
                amount = dec(sheet.cell(rows_at["new"], c).value)
                if amount:
                    quarters.append({**q, "planned_disbursement": amount})
        else:
            issues.append("New loans: the scenario sheet's rows were not recognised; planned disbursement not loaded")
    else:
        issues.append(f"New loans: no '{PIPELINE_RATIO_SHEET}' sheet; planned disbursement not loaded")
    return {"limits": limits, "quarters": quarters, "stated_ratio_pct": stated}


def read_energy_financing(path: Path, portfolio: Dict[str, Any], issues: List[str]) -> None:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    opening = portfolio["quarters"][0]["period_end_ad"] if portfolio["quarters"] else None
    portfolio["bonds"] = read_bonds(book, portfolio.get("bond_holding"), issues)
    portfolio["ratio_quarters"] = read_ratio_inputs(book, opening, issues)
    portfolio["pipeline"] = read_pipeline(book, issues)


def energy_financing_results(portfolio: Dict[str, Any], with_pipeline: bool = False):
    """The ratio as the application will work it out from what is about to be loaded."""
    from backend.app.services import energy_financing as engine

    projected = {q["period_end_ad"]: q["outstanding_parsed"] for q in portfolio["quarters"]}
    additional = engine.cumulative(
        {q["period_end_ad"]: q["planned_disbursement"] for q in portfolio["pipeline"]["quarters"]},
        [q["period_end_ad"] for q in portfolio["ratio_quarters"]]) if with_pipeline else None
    return engine.calculate(
        [engine.QuarterInput(q["period_end_ad"], q["bank_total_loans"], q["required_pct"],
                             q["hydro_outstanding_actual"], q["energy_bond_actual"]) for q in portfolio["ratio_quarters"]],
        projected,
        [engine.Bond(b["amount"], b["investment_date"], b["maturity_date"]) for b in portfolio["bonds"]],
        additional)


def resolve_statuses(borrowers: List[Dict[str, Any]], issues: List[str]) -> None:
    """Rows the legend colours do not cover: a facility with no limit left is settled."""
    for b in borrowers:
        if b["status"] is None:
            if b["limit"] == 0 and b["outstanding"] == 0:
                b["status"] = "settled"
            else:
                b["status"] = "yet to start drawdown" if b["outstanding"] == 0 else "under construction"
                issues.append(f"{b['name']}: no status colour; taken as '{b['status']}' from its balances")


def load_doed() -> Dict[str, List[Dict[str, str]]]:
    by_promoter: Dict[str, List[Dict[str, str]]] = defaultdict(list)
    if DOED_SOURCE.exists():
        with DOED_SOURCE.open(encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle):
                if row["promoter"] and row["capacity_mw"]:
                    by_promoter[name_key(row["promoter"])].append(row)
    return by_promoter


def attach_locations(borrowers: List[Dict[str, Any]]) -> None:
    """Province and district from the DoED licence whose promoter is this borrower and whose size is closest."""
    doed = load_doed()
    for b in borrowers:
        licences = doed.get(name_key(b["name"]), [])
        b["licence"] = min(licences, key=lambda row: abs(Decimal(row["capacity_mw"]) - (b["mw"] or 0))) if licences else None


# ---------------------------------------------------------------------------------------------
# Generation workbook
# ---------------------------------------------------------------------------------------------

def read_generation(path: Path, issues: List[str]) -> Dict[str, Dict[Tuple[int, int], Dict[str, Any]]]:
    """Plant title -> {(BS year, BS month): energy in kWh, revenue and rate as the sheet gives them}."""
    import openpyxl

    plants: Dict[str, Dict[Tuple[int, int], Dict[str, Any]]] = {}
    for sheet in openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        title = next((" ".join(str(r[0]).split()) for r in rows if r and r[0]), sheet.title)
        months: Dict[Tuple[int, int], Dict[str, Any]] = {}
        columns: Dict[str, int] = {}
        block: Optional[Tuple[int, bool]] = None  # (first BS year of the block, block is a fiscal year)
        for number, row in enumerate(rows, start=1):
            text = " ".join(str(c) for c in row if c is not None)
            heading = re.search(r"For\s+(Fiscal\s+)?Year\s+(\d{4})", text, re.I)
            if heading:
                block = (int(heading.group(2)), bool(heading.group(1)))
                continue
            headers = [str(c).strip().lower() if c else "" for c in row]
            if any("month" in h for h in headers):
                columns = {}
                for index, h in enumerate(headers):
                    key = ("month" if "month" in h else "rate" if h.startswith("ppa rate") else
                           "revenue" if "generation revenue" in h else
                           "generated" if "energy generated" in h or "generation energy" in h else
                           "contract" if h.startswith("contract energy") else None)
                    if key:
                        columns[key] = index
                continue
            if not columns or len(row) <= columns["month"]:
                continue
            label = row[columns["month"]]
            month = bs_month(label) if label else None
            if month is None:
                if label and not str(label).strip().lower().startswith("total"):
                    issues.append(f"{title} row {number}: month label {label!r} not understood; skipped")
                continue

            def cell(key: str) -> Optional[Decimal]:
                return dec(row[columns[key]]) if key in columns and columns[key] < len(row) else None

            in_label = re.search(r"\d{4}", str(label))
            if in_label:
                year = int(in_label.group())
            elif block:
                # A fiscal year runs Shrawan to Asadh: its last three months fall in the next BS year
                year = block[0] + (1 if block[1] and month <= 3 else 0)
            else:
                issues.append(f"{title} row {number}: no year for {label!r}; skipped")
                continue
            generated = cell("generated")
            if generated is None:
                continue  # a month the sheet lists but has not filled in yet
            if (year, month) in months:
                issues.append(f"{title}: {year}-{month:02d} appears twice; the later row is used")
            months[(year, month)] = {"contract": cell("contract"), "generated": generated,
                                     "revenue": cell("revenue"), "rate": cell("rate")}
        plants[title] = months
    return plants


def assign_generation(plants: Dict[str, Dict], borrowers: List[Dict[str, Any]], issues: List[str]) -> None:
    """Match each plant sheet to a borrower on the first word of its name. A borrower with several
    plants gets their monthly sum, because the limit is one facility for the borrower as a whole."""
    for b in borrowers:
        b["generation"], b["plants"] = {}, []
    for title, months in plants.items():
        first = name_key(title).split()[0]
        owners = [b for b in borrowers if name_key(b["name"]).split()[0] == first]
        if len(owners) != 1:
            issues.append(f"Generation sheet '{title}': {len(owners)} borrowers match; not loaded")
            continue
        owner = owners[0]
        owner["plants"].append(title)
        for key, values in months.items():
            merged = owner["generation"].setdefault(key, {"contract": None, "generated": Decimal(0), "revenue": None, "rate": None})
            for field in ("contract", "generated", "revenue"):
                if values[field] is not None:
                    merged[field] = (merged[field] or Decimal(0)) + values[field]
            merged["rate"] = values["rate"] or merged["rate"]


def rate_history(generation: Dict[Tuple[int, int], Dict[str, Any]]) -> List[Dict[str, Any]]:
    """The PPA rate per season as dated steps: a new step starts in the month the rate changes."""
    steps: List[Dict[str, Any]] = []
    for season, dry in (("wet", False), ("dry", True)):
        runs: List[Dict[str, Any]] = []
        for (year, month), values in sorted(generation.items()):
            if values["rate"] is None or (month in DRY_MONTHS) != dry:
                continue
            if not runs or runs[-1]["rate"] != values["rate"]:
                start = bs_month_start(year, month)
                if runs:
                    runs[-1]["valid_to_ad"] = start - timedelta(days=1)
                runs.append({"season": season, "rate": values["rate"], "valid_from_ad": start,
                             "valid_from_bs": f"{year}-{month:02d}-01", "valid_to_ad": None})
        steps.extend(runs)
    return steps


# ---------------------------------------------------------------------------------------------
# Report and load
# ---------------------------------------------------------------------------------------------

def report(portfolio: Dict[str, Any], issues: List[str]) -> bool:
    """Print how the parsed rows compare with the totals the workbook states. True when they agree."""
    borrowers, total = portfolio["borrowers"], portfolio["total"]
    sums = {"mw": sum((b["mw"] or 0 for b in borrowers), Decimal(0)),
            "limit": sum((b["limit"] for b in borrowers), Decimal(0)),
            "outstanding": sum((b["outstanding"] for b in borrowers), Decimal(0))}
    print(f"Portfolio ({portfolio['as_of_label']}): {len(borrowers)} borrowers")
    agrees = True
    for field in ("mw", "limit", "outstanding"):
        match = abs(sums[field] - (total.get(field) or 0)) < 1
        agrees = agrees and match
        print(f"  {field:12} parsed {sums[field]:>20,.2f}   workbook {total.get(field) or 0:>20,.2f}   {'ok' if match else 'DIFFERS'}")
    print("  By status (count, MW)                 parsed            workbook")
    for label in STATUSES:
        rows = [b for b in borrowers if b["status"] == label]
        stated = portfolio["stated"].get(label, {"mw": Decimal(0), "count": 0})
        mw = sum((b["mw"] or 0 for b in rows), Decimal(0))
        print(f"    {label:24} {len(rows):>3} {mw:>10,.2f}       {stated['count']:>3} {stated['mw']:>10,.2f}")
    quarters = portfolio["quarters"]
    if quarters:
        projected = [b for b in borrowers if b["projection"]]
        broken = sum(1 for b in projected for i in range(1, len(quarters))
                     if abs(b["projection"][i - 1]["outstanding"] + b["projection"][i]["disbursement"]
                            - b["projection"][i]["repayment"] - b["projection"][i]["outstanding"]) >= 1)
        print(f"  Projection: {len(quarters)} quarters, {quarters[0]['period_end_bs']} to {quarters[-1]['period_end_bs']}, "
              f"for {len(projected)} borrowers")
        print(f"    quarters where outstanding is not previous + disbursement - repayment: {broken}")
        # The borrower rows are what gets loaded. The workbook's total rows are only compared: they are
        # typed or carried forward in places, so a gap is the workbook's and does not stop the load.
        for name in ("disbursement", "repayment", "outstanding"):
            gaps = [(q, (q[f"{name}_stated"] or Decimal(0)) - q[f"{name}_parsed"]) for q in quarters]
            gaps = [(q, gap) for q, gap in gaps if abs(gap) >= 1]
            if not gaps:
                print(f"    {name:12} sum of rows equals the workbook's total row in every quarter")
            else:
                first, widest = gaps[0][0], max(gaps, key=lambda item: abs(item[1]))[1]
                print(f"    {name:12} workbook's total row differs from the sum of its rows in {len(gaps)} quarters, "
                      f"from {first['fiscal_year']} Q{first['quarter']}; widest gap {widest:,.2f}")
        print(f"    outstanding {quarters[0]['outstanding_parsed']:,.2f} -> {quarters[-1]['outstanding_parsed']:,.2f}; "
              f"disbursement {sum(q['disbursement_parsed'] for q in quarters):,.2f}; "
              f"repayment {sum(q['repayment_parsed'] for q in quarters):,.2f}")
        for b in projected:
            if abs(b["projection"][0]["outstanding"] - b["outstanding"]) >= 1:
                issues.append(f"{b['name']}: the projection opens at {b['projection'][0]['outstanding']:,.2f}, "
                              f"not at its principal outstanding of {b['outstanding']:,.2f}")
    bonds, ratio_quarters = portfolio.get("bonds", []), portfolio.get("ratio_quarters", [])
    if bonds or ratio_quarters:
        held = sum((b["amount"] for b in bonds), Decimal(0))
        stated_holding = portfolio.get("bond_holding")
        print(f"  Energy bonds: {len(bonds)} bonds, {held:,.2f}"
              + (f"   projection sheet states {stated_holding:,.2f}" if stated_holding is not None else ""))
        results = [(q, r) for q, r in zip(ratio_quarters, energy_financing_results(portfolio)) if r.ratio_pct is not None]
        if results:
            print(f"  Energy financing ratio: {len(results)} quarters, {results[0][0]['period_end_bs']} to {results[-1][0]['period_end_bs']}")
            apart = [(q, r) for q, r in results
                     if q["stated_ratio_pct"] is not None and abs(q["stated_ratio_pct"] - r.ratio_pct) >= Decimal("0.01")]
            compared = sum(1 for q, _ in results if q["stated_ratio_pct"] is not None)
            print(f"    within 0.01 of the workbook's own ratio in {compared - len(apart)} of the {compared} quarters it states")
            for q, r in apart:
                print(f"      {q['fiscal_year']} Q{q['quarter']}: worked out {r.ratio_pct:.2f}%, workbook {q['stated_ratio_pct']:.2f}%")
            short = [q for q, r in results if r.status == "shortfall"]
            print("    minimum met in every quarter that has one" if not short else
                  f"    below the minimum in {len(short)} quarters, first {short[0]['fiscal_year']} Q{short[0]['quarter']}")
    pipeline = portfolio.get("pipeline") or {"limits": [], "quarters": [], "stated_ratio_pct": {}}
    if pipeline["limits"] or pipeline["quarters"]:
        limit_total = sum((l["new_limit"] for l in pipeline["limits"]), Decimal(0))
        scheduled = sum((q["planned_disbursement"] for q in pipeline["quarters"]), Decimal(0))
        print(f"  New loans planned: {len(pipeline['limits'])} fiscal years, limits {limit_total:,.2f}; "
              f"{scheduled:,.2f} scheduled over {len(pipeline['quarters'])} quarters"
              + (f", {pipeline['quarters'][0]['period_end_bs']} to {pipeline['quarters'][-1]['period_end_bs']}" if pipeline["quarters"] else ""))
        if pipeline["quarters"] and ratio_quarters:
            scenario = [(q, r) for q, r in zip(ratio_quarters, energy_financing_results(portfolio, with_pipeline=True))
                        if r.ratio_pct is not None and not r.is_actual]
            stated = pipeline["stated_ratio_pct"]
            compared = [(q, r, stated[q["period_end_ad"]]) for q, r in scenario if q["period_end_ad"] in stated]
            gaps = [r.ratio_pct - s for _, r, s in compared]
            apart = [g for g in gaps if abs(g) >= Decimal("0.01")]
            print(f"    with new loans: within 0.01 of the workbook's own ratio in {len(compared) - len(apart)} of {len(compared)} projected quarters"
                  + (f"; worked out higher by {min(apart):.2f} to {max(apart):.2f} points in the rest" if apart else ""))
            short = [q for q, r in scenario if r.status == "shortfall"]
            print("    with new loans: minimum met in every projected quarter" if not short else
                  f"    with new loans: below the minimum in {len(short)} quarters, first {short[0]['fiscal_year']} Q{short[0]['quarter']}")
    located = sum(1 for b in borrowers if b["licence"])
    print(f"  Location found in the DoED list for {located} of {len(borrowers)} borrowers")
    print("Generation:")
    for b in borrowers:
        if b["generation"]:
            keys = sorted(b["generation"])
            priced = sum(1 for v in b["generation"].values() if v["revenue"] is not None)
            print(f"  {b['name'][:40]:40} {len(keys):>4} months  {keys[0][0]}-{keys[0][1]:02d} to {keys[-1][0]}-{keys[-1][1]:02d}"
                  f"  with revenue {priced:>4}  rate steps {len(rate_history(b['generation'])):>2}  ({len(b['plants'])} sheet(s))")
    print(f"  Total plant-months: {sum(len(b['generation']) for b in borrowers)}")
    if issues:
        print(f"Notes ({len(issues)}):")
        for issue in issues:
            print(f"  - {issue}")
    return agrees


async def load(portfolio: Dict[str, Any], projection_file: str, generation_file: str) -> Dict[str, int]:
    from sqlalchemy import delete
    from sqlalchemy.ext.asyncio import AsyncSession

    from backend.app.database import engine
    from backend.app.models.financial import (
        EnergyBond, EnergyFinancingQuarter, LoanAccount, LoanProjectionQuarter, NewLoanDisbursementQuarter, NewLoanLimit,
    )
    from backend.app.models.operations import EnergyGenerationData, NEAPPARate
    from backend.app.models.project import Project
    from backend.app.services.risk_service import bs_string
    from backend.scripts.seed_realistic_data import DELETE_ORDER, ID_NAMESPACE, PROVINCES

    if engine.url.host not in ("localhost", "127.0.0.1", "::1"):
        raise SystemExit(f"Refusing to load customer data into {engine.url.host}: local databases only.")

    counts = {"projects": 0, "loans": 0, "projection": 0, "generation": 0, "rates": 0,
              "bonds": len(portfolio["bonds"]), "ratio_quarters": len(portfolio["ratio_quarters"]),
              "new_limits": len(portfolio["pipeline"]["limits"]), "new_loan_quarters": len(portfolio["pipeline"]["quarters"])}
    async with AsyncSession(engine) as session:
        for model in DELETE_ORDER:
            await session.execute(delete(model))
        for bond in portfolio["bonds"]:
            session.add(EnergyBond(
                id=uuid.uuid4(), name=bond["name"], amount=money(bond["amount"]), yield_pct=bond["yield_pct"],
                investment_date_ad=bond["investment_date"], investment_date_bs=bs_string(bond["investment_date"]),
                maturity_date_ad=bond["maturity_date"], maturity_date_bs=bs_string(bond["maturity_date"]),
                data_provenance=PROVENANCE, source_reference=f"{projection_file} / {BOND_SHEET}"[:255]))
        for q in portfolio["ratio_quarters"]:
            session.add(EnergyFinancingQuarter(
                id=uuid.uuid4(), fiscal_year=q["fiscal_year"], quarter=q["quarter"],
                period_end_ad=q["period_end_ad"], period_end_bs=q["period_end_bs"],
                bank_total_loans=money(q["bank_total_loans"]), required_pct=q["required_pct"],
                hydro_outstanding_actual=money(q["hydro_outstanding_actual"]),
                energy_bond_actual=money(q["energy_bond_actual"]),
                data_provenance=PROVENANCE, source_reference=f"{projection_file} / {RATIO_SHEET}"[:255]))

        for limit in portfolio["pipeline"]["limits"]:
            session.add(NewLoanLimit(
                id=uuid.uuid4(), fiscal_year=limit["fiscal_year"], new_limit=money(limit["new_limit"]),
                drawdown_pct=limit["drawdown_pct"], data_provenance=PROVENANCE,
                source_reference=f"{projection_file} / {PIPELINE_SHEET}"[:255]))
        for q in portfolio["pipeline"]["quarters"]:
            session.add(NewLoanDisbursementQuarter(
                id=uuid.uuid4(), fiscal_year=q["fiscal_year"], quarter=q["quarter"], period_end_ad=q["period_end_ad"],
                period_end_bs=q["period_end_bs"], planned_disbursement=money(q["planned_disbursement"]),
                data_provenance=PROVENANCE, source_reference=f"{projection_file} / {PIPELINE_RATIO_SHEET}"[:255]))

        for index, b in enumerate(portfolio["borrowers"], start=1):
            pipeline_status, stage = STATUSES[b["status"]]
            licence = b["licence"] or {}
            province, code = PROVINCES.get(licence.get("province"), (None, "NP"))
            project = Project(
                id=uuid.uuid5(ID_NAMESPACE, f"bank:{b['name']}"), project_code=f"HPM-{code}-{index:04d}",
                name_en=b["name"], name_np=b["name"], province=province,
                district=(licence.get("district") or "").strip().title() or None,
                local_level=(licence.get("municipality") or "").strip() or None,
                installed_capacity_mw=b["mw"] or Decimal(0), project_stage=stage, pipeline_status=pipeline_status)
            session.add(project)
            await session.flush()
            counts["projects"] += 1

            if b["limit"] > 0 or b["outstanding"] > 0:
                session.add(LoanAccount(
                    id=uuid.uuid4(), project_id=project.id, finacle_account_id=f"XLS-{project.project_code}",
                    customer_cif=b["cif"], facility_type="term_loan", sanctioned_amount=money(b["limit"]),
                    disbursed_amount=money(b["limit"] - b["remaining"]), outstanding_principal=money(b["outstanding"]),
                    sync_status="success", data_provenance=PROVENANCE,
                    source_reference=f"{projection_file} / {PORTFOLIO_SHEET} row {b['row']}"[:255]))
                counts["loans"] += 1

            for quarter, figures in zip(portfolio["quarters"], b["projection"]):
                session.add(LoanProjectionQuarter(
                    id=uuid.uuid4(), project_id=project.id, fiscal_year=quarter["fiscal_year"],
                    quarter=quarter["quarter"], period_end_ad=quarter["period_end_ad"],
                    period_end_bs=quarter["period_end_bs"], is_opening=quarter["is_opening"],
                    projected_disbursement=money(figures["disbursement"]),
                    projected_repayment=money(figures["repayment"]),
                    projected_outstanding=money(figures["outstanding"]), data_provenance=PROVENANCE,
                    source_reference=f"{projection_file} / {PORTFOLIO_SHEET}"[:255]))
                counts["projection"] += 1

            for (year, month), values in sorted(b["generation"].items()):
                session.add(EnergyGenerationData(
                    id=uuid.uuid4(), project_id=project.id, month_ad=bs_month_start(year, month),
                    month_bs=f"{year}-{month:02d}", season="dry" if month in DRY_MONTHS else "wet",
                    contract_energy_mwh=money(values["contract"] / 1000) if values["contract"] is not None else None,
                    actual_energy_mwh=money(values["generated"] / 1000), revenue_npr=money(values["revenue"]),
                    data_provenance=PROVENANCE, source_reference=f"{generation_file} / {', '.join(b['plants'])}"[:255]))
                counts["generation"] += 1
            steps = rate_history(b["generation"])
            for step in steps:
                current = step["valid_to_ad"] is None
                session.add(NEAPPARate(
                    id=uuid.uuid4(), project_id=project.id, season=step["season"],
                    rate_per_mwh_npr=step["rate"] * 1000, valid_from_ad=step["valid_from_ad"],
                    valid_from_bs=step["valid_from_bs"], valid_to_ad=step["valid_to_ad"], is_current=current,
                    data_provenance=PROVENANCE, source_reference=generation_file[:255]))
                counts["rates"] += 1
        await session.commit()
    return counts


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--projection", type=Path, required=True, help="the loan projection workbook")
    parser.add_argument("--generation", type=Path, required=True, help="the generation details workbook")
    parser.add_argument("--apply", action="store_true", help="replace the local portfolio with the workbook's")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    issues: List[str] = []
    portfolio = read_portfolio(args.projection, issues)
    resolve_statuses(portfolio["borrowers"], issues)
    attach_locations(portfolio["borrowers"])
    read_energy_financing(args.projection, portfolio, issues)
    assign_generation(read_generation(args.generation, issues), portfolio["borrowers"], issues)
    agrees = report(portfolio, issues)

    if not args.apply:
        print("\nDry run: nothing written. Add --apply to replace the local portfolio.")
        return 0
    if not agrees:
        print("\nNot loaded: the parsed totals do not match the workbook's own totals.")
        return 1
    counts = asyncio.run(load(portfolio, args.projection.name, args.generation.name))
    print(f"\nLoaded {counts['projects']} projects, {counts['loans']} loan accounts, "
          f"{counts['projection']} projection quarters, {counts['bonds']} bonds, {counts['new_limits']} new-loan limits, "
          f"{counts['new_loan_quarters']} new-loan quarters, "
          f"{counts['ratio_quarters']} ratio quarters, {counts['generation']} generation months, "
          f"{counts['rates']} rate steps.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
